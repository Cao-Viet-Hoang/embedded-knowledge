#!/usr/bin/env python
"""Convert a CV Markdown file to a styled DOCX (python-docx), mirroring build_cv_pdf.py.

Variants built from the same Markdown source:
  - <stem>.docx      : with avatar photo (for VN/JP/DE submissions)
  - <stem>_ATS.docx  : no photo (for US/global ATS systems)

Supported Markdown subset: # / ## / ### / #### headings, paragraphs, nested "-" lists,
**bold**, *italic* / _italic_, [links](url), and --- rules.

Usage:
  python build_cv_docx.py                       # default CV, both variants
  python build_cv_docx.py --no-photo            # default CV, ATS only
  python build_cv_docx.py --photo               # default CV, photo only
  python build_cv_docx.py --src ../temp_cv/Foo.md --no-photo --out ../temp_cv/Foo.docx
"""
import argparse
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

DEFAULT_SRC = Path("CV_Cao_Viet_Hoang.md")
DEFAULT_AVATAR_NAME = "CV_picture_avatar.jpg"

FONT = "Calibri"
NAVY = RGBColor(0x1A, 0x3C, 0x66)
BLUE = RGBColor(0x2C, 0x5D, 0x99)
TEXT = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)

PAGE_W_MM, MARGIN_MM = 210, 16
AVATAR_W_MM, AVATAR_H_MM = 28, 36  # 105x135 px in the PDF

INLINE_RE = re.compile(
    r"\*\*(?P<b>.+?)\*\*"
    r"|\[(?P<lt>[^\]]+)\]\((?P<lu>[^)]+)\)"
    r"|(?<![\w*])_(?P<i1>.+?)_(?![\w*])"
    r"|(?<![\w*])\*(?P<i2>.+?)\*(?![\w*])"
)


def display_name_from(src):
    stem = src.stem
    if stem.startswith("CV_"):
        stem = stem[3:]
    if stem.endswith("_CV"):
        stem = stem[:-3]
    return stem.replace("_", " ").strip() or src.stem


# ---------- low-level docx helpers ----------

def style_run(run, size=None, bold=False, italic=False, color=TEXT):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_hyperlink(par, text, url, size, bold=False):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = par.add_run(text)
    style_run(run, size=size, bold=bold, color=BLUE)
    link.append(run._element)
    par._p.append(link)


def add_inline(par, text, size, color=TEXT, bold=False, italic=False, bold_color=NAVY):
    """Render inline Markdown into runs; nests bold/italic/link."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            style_run(par.add_run(text[pos:m.start()]), size, bold, italic, color)
        if m.group("b") is not None:
            add_inline(par, m.group("b"), size, color=bold_color if not bold else color,
                       bold=True, italic=italic, bold_color=bold_color)
        elif m.group("lt") is not None:
            add_hyperlink(par, m.group("lt"), m.group("lu"), size, bold)
        else:
            inner = m.group("i1") if m.group("i1") is not None else m.group("i2")
            add_inline(par, inner, size, color=GREY, bold=bold, italic=True,
                       bold_color=bold_color)
        pos = m.end()
    if pos < len(text):
        style_run(par.add_run(text[pos:]), size, bold, italic, color)


def set_spacing(par, before=0, after=0, line=1.15, keep_next=False):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.keep_together = True


def bottom_border(par, color="1A3C66", size=12):
    pPr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", "2"), ("color", color)):
        bottom.set(qn(f"w:{k}"), v)
    borders.append(bottom)
    pPr.append(borders)


def hide_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def set_list_indent(par, level):
    left = 5 + level * 5  # mm
    par.paragraph_format.left_indent = Mm(left)
    par.paragraph_format.first_line_indent = Mm(-4)


# ---------- markdown -> blocks ----------

def parse_blocks(text):
    """Yield (kind, level, text): kind in h/p/li/hr; level = heading depth or list nesting."""
    para = []

    def flush():
        if para:
            yield ("p", 0, " ".join(s.strip() for s in para))
            para.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            yield from flush()
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        li = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if h:
            yield from flush()
            yield ("h", len(h.group(1)), h.group(2).strip())
        elif re.match(r"^\s*(-{3,}|\*{3,})\s*$", line):
            yield from flush()
            yield ("hr", 0, "")
        elif li:
            yield from flush()
            indent = len(li.group(1).replace("\t", "    "))
            yield ("li", indent // 2 if indent < 4 else indent // 4, li.group(2).strip())
        else:
            para.append(line)
    yield from flush()


# ---------- builders ----------

def render_block(container, kind, level, text, state):
    """Add one block to a Document or table cell."""
    if kind == "h":
        if level == 1:
            par = container.add_paragraph()
            add_inline(par, text, 24, color=NAVY, bold=True)
            set_spacing(par, 0, 3, 1.0, keep_next=True)
            state["after_h1"] = True
            return
        state["after_h1"] = False
        par = container.add_paragraph()
        if level == 2:
            add_inline(par, text.upper(), 14, color=NAVY, bold=True)
            set_spacing(par, 14, 6, 1.0, keep_next=True)
            bottom_border(par)
        elif level == 3:
            add_inline(par, text, 11, color=RGBColor(0x33, 0x33, 0x33), bold=True)
            set_spacing(par, 10, 2, 1.1, keep_next=True)
        else:
            add_inline(par, text, 10.5, color=BLUE, bold=True)
            set_spacing(par, 8, 2, 1.1, keep_next=True)
    elif kind == "p":
        par = container.add_paragraph()
        if state.get("after_h1"):  # role tagline
            add_inline(par, text, 11.5, color=BLUE, bold=True, bold_color=BLUE)
            set_spacing(par, 0, 4, 1.1)
            state["after_h1"] = False
        else:
            add_inline(par, text, 10.5)
            set_spacing(par, 3, 3)
    elif kind == "li":
        state["after_h1"] = False
        par = container.add_paragraph()
        add_inline(par, "•\t", 10.5)
        add_inline(par, text, 10.5)
        set_spacing(par, 2, 2)
        set_list_indent(par, level)
        par.paragraph_format.tab_stops.add_tab_stop(Mm(5 + level * 5))
    elif kind == "hr":
        par = container.add_paragraph()
        set_spacing(par, 2, 4, 1.0)
        bottom_border(par, color="CCCCCC", size=6)


def avatar_stream(avatar):
    """Center-crop the avatar to the target aspect (like CSS object-fit: cover)."""
    try:
        from PIL import Image
    except ImportError:
        return str(avatar)
    img = Image.open(avatar).convert("RGB")
    target = AVATAR_W_MM / AVATAR_H_MM
    w, h = img.size
    if w / h > target:
        new_w = int(h * target)
        img = img.crop(((w - new_w) // 2, 0, (w - new_w) // 2 + new_w, h))
    else:
        new_h = int(w / target)
        img = img.crop((0, (h - new_h) // 2, w, (h - new_h) // 2 + new_h))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=92)
    buf.seek(0)
    return buf


def build_docx(src, out, name, include_photo=True, avatar=None):
    blocks = list(parse_blocks(src.read_text(encoding="utf-8")))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(PAGE_W_MM), Mm(297)
    sec.left_margin = sec.right_margin = Mm(MARGIN_MM)
    sec.top_margin = sec.bottom_margin = Mm(MARGIN_MM)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)

    doc.core_properties.title = f"{name} — CV"
    doc.core_properties.author = name

    state = {}
    use_photo = include_photo and avatar is not None and avatar.exists()

    # Header = everything before the first H2.
    first_h2 = next((i for i, b in enumerate(blocks) if b[0] == "h" and b[1] == 2), len(blocks))
    head, body = blocks[:first_h2], blocks[first_h2:]

    if use_photo:
        text_w = PAGE_W_MM - 2 * MARGIN_MM
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        hide_table_borders(table)
        left, right = table.rows[0].cells
        left.width, right.width = Mm(text_w - AVATAR_W_MM - 6), Mm(AVATAR_W_MM + 6)
        for blk in head:
            render_block(left, *blk, state)
        left._tc.remove(left.paragraphs[0]._p)  # drop the empty default paragraph
        rpar = right.paragraphs[0]
        rpar.alignment = 2  # right
        rpar.add_run().add_picture(avatar_stream(avatar), width=Mm(AVATAR_W_MM),
                                   height=Mm(AVATAR_H_MM))
    else:
        for blk in head:
            render_block(doc, *blk, state)

    for blk in body:
        render_block(doc, *blk, state)

    # python-docx's default template starts with one empty paragraph in some cases.
    first = doc.paragraphs[0] if doc.paragraphs else None
    if first is not None and not first.text.strip() and not first._p.xpath(".//w:drawing"):
        if doc.element.body.index(first._p) == 0:
            doc.element.body.remove(first._p)

    doc.save(out)
    print(f"Wrote {out}")


def main():
    parser = argparse.ArgumentParser(description="Build the CV DOCX file(s).")
    parser.add_argument("--src", type=Path, default=DEFAULT_SRC,
                        help="source CV Markdown file (default: CV_Cao_Viet_Hoang.md)")
    parser.add_argument("--out-dir", type=Path, default=None,
                        help="directory for outputs (default: alongside --src)")
    parser.add_argument("--out", type=Path, default=None,
                        help="explicit output DOCX path; requires --photo or --no-photo")
    parser.add_argument("--avatar", type=Path, default=None,
                        help="avatar image (default: CV_picture_avatar.jpg next to --src)")
    parser.add_argument("--name", default=None,
                        help="display name for document properties (default: derived from --src)")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--no-photo", action="store_true",
                       help="build only the no-photo ATS variant")
    group.add_argument("--photo", action="store_true",
                       help="build only the photo variant")
    args = parser.parse_args()

    src = args.src
    if not src.exists():
        raise FileNotFoundError(f"Source markdown not found: {src}")

    name = args.name or display_name_from(src)
    avatar = args.avatar or (src.parent / DEFAULT_AVATAR_NAME)

    if args.out is not None:
        if not (args.photo or args.no_photo):
            parser.error("--out requires --photo or --no-photo (a single variant)")
        args.out.parent.mkdir(parents=True, exist_ok=True)
        jobs = [(bool(args.photo), args.out)]
    else:
        out_dir = args.out_dir or src.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        photo_out = out_dir / f"{src.stem}.docx"
        ats_out = out_dir / f"{src.stem}_ATS.docx"
        if args.no_photo:
            jobs = [(False, ats_out)]
        elif args.photo:
            jobs = [(True, photo_out)]
        else:
            jobs = [(True, photo_out), (False, ats_out)]

    for include_photo, out in jobs:
        build_docx(src, out, name, include_photo=include_photo, avatar=avatar)


if __name__ == "__main__":
    main()
